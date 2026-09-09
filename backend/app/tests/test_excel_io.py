from io import BytesIO

from openpyxl import load_workbook

from app import ledger_config, models
from app.services import excel_io


def test_template_headers_and_example_row():
    defn = ledger_config.get_ledger("dibao")
    wb = load_workbook(BytesIO(excel_io.build_template(defn)))
    ws = wb.active
    headers = [str(c.value).replace(" *", "").replace("*", "").strip() for c in ws[1]]
    assert "户主姓名" in headers and "户主身份证" in headers and "月保障金" in headers
    assert str(ws.cell(2, 1).value).startswith("[示例]")
    assert ws.cell(3, 1).value is None  # 数据从第 3 行开始


def test_parse_and_apply_import(client_db):
    defn = ledger_config.get_ledger("dibao")
    wb = load_workbook(BytesIO(excel_io.build_template(defn)))
    ws = wb.active
    headers = {str(c.value).replace(" *", "").replace("*", "").strip(): i
               for i, c in enumerate(ws[1], start=1) if c.value}
    row_no = 3
    ws.cell(row_no, headers["户主姓名"]).value = "王建国"
    ws.cell(row_no, headers["户主身份证"]).value = "110101195803041234"
    ws.cell(row_no, headers["保障人数"]).value = 3
    ws.cell(row_no, headers["月保障金"]).value = 930
    buf = BytesIO()
    wb.save(buf)
    rows, errors = excel_io.parse_import(defn, buf.getvalue())
    assert errors == []
    assert rows[0]["hz_idcard"] == "110101195803041234"
    s = client_db.SessionLocal()
    result = excel_io.apply_import(s, defn, rows)
    s.commit()
    s.close()
    assert result["added"] == 1
    s = client_db.SessionLocal()
    assert s.query(models.Household).count() == 1
    assert s.query(models.HouseholdTag).filter_by(tag_type="dibao").count() == 1
    s.close()


def test_import_reports_bad_idcard_row():
    defn = ledger_config.get_ledger("dibao")
    wb = load_workbook(BytesIO(excel_io.build_template(defn)))
    ws = wb.active
    headers = {str(c.value).replace(" *", "").replace("*", "").strip(): i
               for i, c in enumerate(ws[1], start=1) if c.value}
    ws.cell(3, headers["户主姓名"]).value = "张三"
    ws.cell(3, headers["户主身份证"]).value = "123"
    buf = BytesIO()
    wb.save(buf)
    rows, errors = excel_io.parse_import(defn, buf.getvalue())
    assert rows == []
    assert any("身份证" in e["msg"] for e in errors)
    assert errors[0]["row"] == 3
