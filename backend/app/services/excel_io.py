import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app import db, ledger_config, models
from app.services import records

IDCARD_RE = re.compile(r"^\d{17}[\dXx]$")
EXAMPLE_PREFIX = "[示例] "
H_HEADERS = [("hz_name", "户主姓名", True), ("hz_idcard", "户主身份证", True),
             ("phone", "联系电话", False), ("address", "住址", False)]
P_HEADERS = [("name", "姓名", True), ("idcard", "身份证", True), ("gender", "性别", False),
             ("birth", "出生日期", False), ("relation", "与户主关系", False),
             ("education", "文化程度", False), ("health", "健康状况", False),
             ("skill", "劳动技能", False)]
TAG_H = [("status", "状态", False), ("start_date", "纳入时间", False), ("end_date", "退出时间", False)]
TAG_P = [("period", "年度", False), ("status", "状态", False),
         ("start_date", "开始时间", False), ("end_date", "结束时间", False)]


def _columns(defn):
    if defn.tag_type is None:
        return P_HEADERS + [(f.key, f.label, f.required) for f in defn.fields]
    if defn.scope == "household":
        return H_HEADERS + TAG_H + [(f.key, f.label, f.required) for f in defn.fields]
    return P_HEADERS + TAG_P + [(f.key, f.label, f.required) for f in defn.fields]


def _style_header(ws, ncols):
    fill = PatternFill("solid", fgColor="B01B2E")
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill


def build_template(defn) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = defn.name
    cols = _columns(defn)
    for i, (key, label, required) in enumerate(cols, start=1):
        ws.cell(1, i, (label + " *" if required else label))
    example = {}
    for key, label, required in cols:
        if key in ("hz_name", "name"):
            example[key] = "张伟"
        elif key in ("hz_idcard", "idcard"):
            example[key] = "110101197001011234"
        elif key == "gender":
            example[key] = "男"
        elif key == "relation":
            example[key] = "户主"
        elif key in ("start_date", "birth"):
            example[key] = "1970-01-01"
        elif key == "period":
            example[key] = "2026"
    for i, (key, label, required) in enumerate(cols, start=1):
        v = example.get(key)
        if v is not None:
            ws.cell(2, i, EXAMPLE_PREFIX + str(v))
    ws.cell(2, len(cols) + 1, EXAMPLE_PREFIX + "← 示例行，导入时自动跳过，可删除")
    for i, (key, label, required) in enumerate(cols, start=1):
        f = next((x for x in defn.fields if x.key == key), None)
        if f and f.options:
            dv = DataValidation(type="list", formula1='"' + ",".join(f.options) + '"')
            ws.add_data_validation(dv)
            dv.add(ws.cell(3, i).coordinate + ":" + get_column_letter(i) + "1000")
    _style_header(ws, len(cols))
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def parse_import(defn, file_bytes):
    wb = load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    headers = {}
    for c in range(1, ws.max_column + 1):
        v = _parse_cell(ws.cell(1, c).value).replace(" *", "").replace("*", "").strip()
        if v:
            headers[c] = v
    coldefs = _columns(defn)
    label_map = {label: key for key, label, req in coldefs}
    rows, errors, seen = [], [], set()
    for r in range(2, ws.max_row + 1):
        first = _parse_cell(ws.cell(r, 1).value)
        if not first or first.startswith(EXAMPLE_PREFIX):
            continue  # 空行与 [示例] 行跳过
        data = {}
        for c, label in headers.items():
            key = label_map.get(label)
            if key:
                data[key] = _parse_cell(ws.cell(r, c).value)
        idcard = data.get("hz_idcard") or data.get("idcard") or ""
        if idcard and not IDCARD_RE.match(idcard):
            errors.append({"row": r, "msg": f"身份证号“{idcard}”位数不正确"})
            continue
        for key, label, required in coldefs:
            if required and not data.get(key):
                errors.append({"row": r, "msg": f"必填项 {label} 缺失"})
                break
        else:
            for f in defn.fields:
                if f.options and data.get(f.key) and data[f.key] not in f.options:
                    errors.append({"row": r, "msg": f"{f.label} 取值“{data[f.key]}”不在可选范围"})
                    break
            else:
                if idcard in seen:
                    errors.append({"row": r, "msg": f"身份证 {idcard} 在文件内重复"})
                    continue
                seen.add(idcard)
                rows.append(data)
    return rows, errors


def _apply_row(s, defn, data):
    extra_keys = {f.key for f in defn.fields}
    extra = {k: v for k, v in data.items() if k in extra_keys and v not in ("", None)}
    for k in ("member_num", "monthly_amount", "payment_amount", "monthly_income", "monthly_pension"):
        if k in extra and extra[k] not in ("", None):
            try:
                extra[k] = float(extra[k])
            except ValueError:
                extra[k] = extra[k]
    from app.services.records import find_or_create_household, find_or_create_person
    if defn.tag_type is None:
        # 居民信息：按身份证存在与否区分新增/更新
        idcard = (data.get("idcard") or "").strip()
        p = s.query(models.Person).filter_by(idcard=idcard).first()
        if p:
            for k in records.BASE_PERSON_FIELDS:
                if k in data and data[k] not in ("", None):
                    setattr(p, k, data[k])
            return "updated"
        find_or_create_person(s, data)
        return "added"
    if defn.scope == "household":
        h = find_or_create_household(s, data)
        tag = models.HouseholdTag(household_id=h.id, tag_type=defn.tag_type,
                                  status=data.get("status") or "享受中",
                                  start_date=_d(data.get("start_date")),
                                  end_date=_d(data.get("end_date")), extra=extra,
                                  remark=data.get("remark") or "")
    else:
        p = find_or_create_person(s, data)
        tag = models.PersonTag(person_id=p.id, tag_type=defn.tag_type,
                               period=data.get("period") or "",
                               start_date=_d(data.get("start_date")),
                               end_date=_d(data.get("end_date")), extra=extra,
                               remark=data.get("remark") or "")
    s.add(tag)
    return "added"


def _d(v):
    if not v:
        return None
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def apply_import(s, defn, rows):
    added = updated = 0
    for data in rows:
        kind = _apply_row(s, defn, data)
        if kind == "added":
            added += 1
        else:
            updated += 1
    return {"added": added, "updated": updated}


def export_rows(defn, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = defn.name
    cols = _columns(defn)
    for i, (key, label, required) in enumerate(cols, start=1):
        ws.cell(1, i, label + (" *" if required else ""))
    _style_header(ws, len(cols))
    for r, row in enumerate(rows, start=2):
        for i, (key, label, required) in enumerate(cols, start=1):
            ws.cell(r, i, row.get(key, ""))
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_all() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for defn in ledger_config.LEDGER_LIST:
        ws = wb.create_sheet(defn.name)
        cols = _columns(defn)
        for i, (key, label, required) in enumerate(cols, start=1):
            ws.cell(1, i, label + (" *" if required else ""))
        _style_header(ws, len(cols))
        with db.SessionLocal() as s:
            if defn.tag_type is None:
                objs = s.query(models.Person).all()
            else:
                model = models.HouseholdTag if defn.scope == "household" else models.PersonTag
                objs = s.query(model).filter(model.tag_type == defn.tag_type).all()
            r = 2
            for obj in objs:
                row = {}
                if defn.tag_type is None:
                    row = {k: getattr(obj, k, "") for k, _, _ in cols}
                elif defn.scope == "household":
                    h = s.get(models.Household, obj.household_id)
                    row = {"hz_name": h.hz_name if h else "", "hz_idcard": h.hz_idcard if h else "",
                           "phone": h.phone if h else "", "address": h.address if h else "",
                           "status": obj.status, "start_date": _fmt(obj.start_date),
                           "end_date": _fmt(obj.end_date), **{k: obj.extra.get(k, "") if obj.extra else "" for k in [f.key for f in defn.fields]}}
                else:
                    p = s.get(models.Person, obj.person_id)
                    row = {"name": p.name if p else "", "idcard": p.idcard if p else "",
                           "period": obj.period, "start_date": _fmt(obj.start_date),
                           "end_date": _fmt(obj.end_date), **{k: obj.extra.get(k, "") if obj.extra else "" for k in [f.key for f in defn.fields]}}
                for i, (key, label, required) in enumerate(cols, start=1):
                    ws.cell(r, i, row.get(key, ""))
                r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(v):
    return v.isoformat() if isinstance(v, date) else ("" if v is None else v)
